from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/mac/Documents/Missubibi')
BUILD = ROOT / '.docx-build'
OUT = ROOT / 'documents' / 'mitsubishi-sales-intelligence-system-overview.docx'
LOGO = BUILD / 'mitsubishi-mark.png'

RED = 'E32439'
INK = '202529'
MUTED = '677780'
LINE = 'D9E0E3'
PALE = 'F4F6F7'
PALE_RED = 'FBEDEE'
WHITE = 'FFFFFF'


def make_logo(path: Path):
    im = Image.new('RGBA', (600, 540), (255, 255, 255, 0))
    d = ImageDraw.Draw(im)
    c = (227, 36, 57, 255)
    d.polygon([(300, 0), (400, 180), (300, 360), (200, 180)], fill=c)
    d.polygon([(300, 360), (200, 540), (0, 540), (100, 360)], fill=c)
    d.polygon([(300, 360), (500, 360), (600, 540), (400, 540)], fill=c)
    im.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=LINE, size='6'):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=80, start=130, bottom=80, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_row_cant_split(row):
    """Keep a table row together so a short continuation does not orphan on the next page."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn('w:cantSplit'))
    if cant_split is None:
        cant_split = OxmlElement('w:cantSplit')
        tr_pr.append(cant_split)


def set_run_font(run, name='Aptos', size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(p, space_before=0, space_after=6, line=1.12, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True


def add_text(doc, text, style=None, size=10.5, color=INK, bold=False, italic=False,
             align=None, before=0, after=6, line=1.12, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    style_paragraph(p, before, after, line, keep)
    return p


def add_rich_paragraph(doc, pieces, before=0, after=6, line=1.12):
    p = doc.add_paragraph()
    for text, opts in pieces:
        r = p.add_run(text)
        set_run_font(r, size=opts.get('size', 10.5), color=opts.get('color', INK),
                     bold=opts.get('bold', False), italic=opts.get('italic', False),
                     name=opts.get('name', 'Aptos'))
    style_paragraph(p, before, after, line)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    style_paragraph(p, space_after=3, line=1.08)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    style_paragraph(p, space_after=3, line=1.08)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), PALE)
    p_pr.append(shd)
    r = p.add_run(text)
    set_run_font(r, name='Courier New', size=9.2, color=INK)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_row_cant_split(table.rows[0])
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        if widths:
            hdr[i].width = Inches(widths[i])
        set_cell_shading(hdr[i], INK)
        set_cell_border(hdr[i])
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, space_after=0, line=1.0)
        r = p.add_run(h)
        set_run_font(r, size=font_size, color=WHITE, bold=True)
    for ridx, row in enumerate(rows):
        table_row = table.add_row()
        set_row_cant_split(table_row)
        cells = table_row.cells
        for i, value in enumerate(row):
            if widths:
                cells[i].width = Inches(widths[i])
            set_cell_shading(cells[i], WHITE if ridx % 2 == 0 else PALE)
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            style_paragraph(p, space_after=0, line=1.05)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_section_heading(doc, text, level=1):
    style = 'Heading 1' if level == 1 else 'Heading 2'
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, size=17 if level == 1 else 13.2, color=INK, bold=True)
    style_paragraph(p, space_before=14 if level == 1 else 10, space_after=6, line=1.0, keep=True)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run('Page ')
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('MITSUBISHI ELECTRIC  ·  Sales Intelligence Platform')
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    add_page_number(footer.add_paragraph())


def add_page_break(doc):
    doc.add_page_break()


def main():
    BUILD.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    make_logo(LOGO)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    add_footer(sec)

    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    styles['Normal'].font.size = Pt(10.5)
    for name, size in [('Title', 28), ('Heading 1', 17), ('Heading 2', 13.2)]:
        styles[name].font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
        styles[name]._element.rPr.rFonts.set(qn('w:ascii'), styles[name].font.name)
        styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), styles[name].font.name)
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(INK)
        styles[name].font.bold = True

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(LOGO), width=Inches(1.65))
    add_text(doc, 'Mitsubishi Electric Sales Intelligence System', style='Title', size=28,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=5, line=1.0)
    add_text(doc, 'System overview and operating guide', size=14, color=MUTED,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    add_text(doc, 'Changes for the Better', size=13, italic=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
    add_text(doc, 'A practical guide to forecasting, demand classification, replenishment planning, data lineage, and teammate setup.',
             size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.2)
    add_page_break(doc)

    # Opening and contents
    add_section_heading(doc, 'Purpose and scope', 1)
    add_text(doc, 'Missubibi is a Mitsubishi Electric Thailand sales intelligence and demand-planning platform. It connects verified Mitsubishi sales history in MySQL to forecasting, item-level demand classification, and lead-time inventory planning. This document explains what the software does, what controls are available, how the data moves through the system, and how teammates can start the current version with the matching database snapshot.', after=8)
    add_text(doc, 'The central planning question is simple: what has sold, what is likely to sell next, which items need attention, and when should replenishment be ordered? The system keeps the evidence and the planning assumptions visible together so a forecast can be explained rather than treated as a black box.', after=10)
    add_section_heading(doc, 'Contents', 1)
    for item in [
        '1  System at a glance',
        '2  Forecast workspace',
        '3  Model performance',
        '4  Demand classifier',
        '5  Lead-time inventory',
        '6  Data and coverage',
        '7  CSV imports and data quality',
        '8  Architecture and operations',
        '9  Teammate setup and database synchronization',
        '10  Interpretation rules and operational boundaries',
    ]:
        add_bullet(doc, item)
    add_page_break(doc)

    # 1
    add_section_heading(doc, '1 System at a glance', 1)
    add_text(doc, 'The dashboard is organized as one planning loop. A user can begin with the data check, move into a forecast for a particular slice of the business, inspect the evidence behind that forecast, prioritize items with the classifiers, and finish with a replenishment queue.', after=8)
    add_table(doc, ['Section', 'What it answers', 'Main user action'], [
        ['Forecast workspace', 'What demand or revenue is likely for a selected business slice?', 'Set scope, measure, frequency, calendar, horizon, and model; then run the forecast.'],
        ['Model performance', 'Which forecasting engine has the strongest historical evidence?', 'Compare accuracy, WAPE, MASE, folds, weights, status, and forecast shape.'],
        ['Demand classifier', 'Which items are likely to be high demand or sold in the next eligible month?', 'Train or refresh the September models, check items, and review borderline probabilities.'],
        ['Lead-time inventory', 'When should stock be ordered and how much is suggested?', 'Set lead time and coverage assumptions, edit on-hand stock, and review the queue.'],
        ['Data and coverage', 'What data is connected and how trustworthy is the current planning view?', 'Inspect dimensions, date coverage, source rows, duplicates, missing prices, and fingerprints.'],
    ], widths=[1.45, 2.7, 2.7], font_size=9.1)
    add_text(doc, 'The backend uses one canonical SQL adapter for forecasts, classifiers, inventory planning, and data coverage. That shared adapter is what lets a user trace a result back to the same source snapshot and filtering rules.', after=8)

    # 2
    add_section_heading(doc, '2 Forecast workspace', 1)
    add_text(doc, 'The Forecast workspace is the main analysis screen. It lets a planner move from a company-wide view to a specific item, customer, product category, industry, or vendor view, then choose the time and value representation that matches the business question.', after=8)
    add_section_heading(doc, 'Available controls', 2)
    add_table(doc, ['Control', 'Options'], [
        ['Scope', 'Company, item / SKU, product category, customer, industry, or vendor.'],
        ['Combined filters', 'Item, customer, industry, and vendor filters can be applied together.'],
        ['Measure', 'Revenue in THB or sales volume in units.'],
        ['Frequency', 'Daily, weekly, monthly, quarterly, or yearly.'],
        ['Calendar', 'Calendar year or fiscal year beginning in April.'],
        ['Horizon', '1 to 120 periods; yearly and quarterly horizons have tighter limits.'],
        ['Model', 'Trend, SARIMA, Random Forest, LSTM, XGBoost, or Ensemble.'],
    ], widths=[1.55, 5.3])
    add_section_heading(doc, 'How to use the forecast', 2)
    add_number(doc, 'Choose the scope and select the item, customer, product, industry, or vendor value when the scope is not the full company.')
    add_number(doc, 'Choose revenue or units, then select the daily, weekly, monthly, quarterly, or yearly view.')
    add_number(doc, 'Choose the calendar convention and set the horizon with the slider.')
    add_number(doc, 'Select a model and click Run forecast. The request runs as a background job and the page polls for the completed result.')
    add_number(doc, 'Read actual history, model forecast, prediction range, accuracy, and totals together.')
    add_number(doc, 'Drag forecast points or edit the period table when you want to test a higher or lower planning scenario.')
    add_rich_paragraph(doc, [('Scenario meaning: ', {'bold': True, 'color': RED}), ('a scenario is an editable planning version of the model forecast. It changes the planning values and scenario fit indicator, but it never rewrites the historical validation score.', {})], before=4, after=8)
    add_text(doc, 'Forecast output can be exported to CSV with the model forecast, planning scenario, lower prediction range, and upper prediction range for each future period.', after=8)

    # 3
    add_section_heading(doc, '3 Model performance', 1)
    add_text(doc, 'The Model performance section compares the six forecast engines using chronological held-out validation against the repaired Mitsubishi SQL snapshot. The displayed scores are calculated evidence, not default ratings and not values adjusted by the planning scenario.', after=8)
    add_table(doc, ['Model', 'Role', 'Held-out accuracy'], [
        ['Trend', 'Interpretable baseline', '78.72%'],
        ['SARIMA', 'Seasonal statistical model', '83.48%'],
        ['Random Forest', 'Nonlinear demand patterns', '84.42%'],
        ['LSTM', 'Sequential neural network', '86.41%'],
        ['XGBoost', 'Gradient-boosted trees', '75.87%'],
        ['Ensemble', 'Validation-weighted blend', '85.20%'],
    ], widths=[1.55, 3.7, 1.6], font_size=9.5)
    add_section_heading(doc, 'What the comparison table shows', 2)
    for item in [
        'Accuracy: max(0, 100 minus validation WAPE).',
        'WAPE: total absolute forecast error divided by total absolute actual demand.',
        'CV MASE: error scaled against a seasonal-naive reference.',
        'Folds: the number of expanding chronological validation windows.',
        'Weight: the contribution used when an ensemble forecast is built.',
        'Outlook: a sparkline showing the model forecast shape for the selected series.',
    ]:
        add_bullet(doc, item)
    add_text(doc, 'Two independent validation runs on the same 265,135-row canonical SQL view reproduced the forecast scores to eight decimal places. A score should change when the SQL data, model code, model configuration, output granularity, or validation protocol changes. A page refresh or scenario edit should not change it.', after=8)
    add_text(doc, 'Prediction ranges are useful planning ranges. Weighted ensemble bounds are approximate and are not claimed to be calibrated statistical 90 percent intervals.', color=MUTED, italic=True, after=8)

    # 4
    add_section_heading(doc, '4 Demand classifier', 1)
    add_text(doc, 'The classifier section sits beside the numeric forecast. It helps prioritize item review by answering two different questions: which items are likely to be high demand, and which items are likely to sell or not sell in the next eligible month.', after=8)
    add_table(doc, ['Model', 'Target', 'Held-out accuracy', 'Other diagnostics'], [
        ['Neural network rev9', 'High demand', '96.92%', 'Balanced accuracy 83.73%; baseline 95.81%; AUC about 0.972; up to 5,706 items.'],
        ['Tabular transformer rev8', 'Sold / not sold', '89.61%', 'Balanced accuracy 81.80%; baseline 86.49%; AUC about 0.915; 5,620 items with sufficient history.'],
    ], widths=[1.65, 1.45, 1.35, 2.4], font_size=8.9)
    add_section_heading(doc, 'How to use the classifier', 2)
    for item in [
        'Choose Quick training for up to 12 epochs or Full training for up to 100 epochs.',
        'Search for one item / SKU or leave the field blank to score all eligible items.',
        'Click Check demand to request probability rows from both classifiers.',
        'Use the borderline review queue to focus on probabilities that need human review.',
        'Read held-out accuracy, balanced accuracy, ROC AUC, the last-month baseline, held-out item-month count, and the confusion matrix.',
    ]:
        add_bullet(doc, item)
    add_text(doc, 'Classifier outputs prioritize attention; they do not replace the revenue or unit forecast. If a model export is missing or stale, the system returns an unavailable state instead of fabricating probabilities.', after=8)

    # 5
    add_section_heading(doc, '5 Lead-time inventory', 1)
    add_text(doc, 'Lead-time inventory connects the sales run rate to the date a purchase order needs to be placed and the date stock is expected to arrive ready for use. Demand comes from the canonical Mitsubishi SQL history. Supplier lead time and warehouse on-hand are explicit planner inputs.', after=8)
    add_table(doc, ['Planning input', 'Control range or source', 'Effect'], [
        ['Supplier lead time', '1 to 365 days', 'Determines demand during the supply wait and order-by date.'],
        ['Safety stock', '0 to 12 months', 'Adds a reserve above lead-time demand.'],
        ['Forward coverage', '0.25 to 24 months', 'Sets target demand coverage after arrival.'],
        ['On-hand units', 'Editable per item', 'Reduces the suggested order and determines Covered or Order now.'],
    ], widths=[1.7, 1.9, 3.2], font_size=9.2)
    add_page_break(doc)
    add_section_heading(doc, 'Planning calculations', 2)
    add_code(doc, 'Demand run rate = 70% trailing three-month demand + 30% trailing twelve-month demand\nLead-time demand = demand run rate × supplier lead-time days\nReorder point = lead-time demand + safety stock\nTarget stock = lead-time demand + safety stock + forward coverage\nSuggested order = target stock − on-hand units')
    add_text(doc, 'The replenishment queue shows item, monthly run rate, demand in lead time, safety stock, reorder point, editable on-hand units, suggested order quantity, order-by date, expected arrival date, and status.', after=8)
    add_text(doc, 'The source Mitsubishi files do not contain purchase orders, receipts, supplier lead times, or warehouse balances. The interface keeps those values visible as planning inputs instead of hiding assumptions inside the calculation.', color=MUTED, italic=True, after=8)

    # 6
    add_section_heading(doc, '6 Data and coverage', 1)
    add_text(doc, 'The data section is the provenance check before a forecast is interpreted. It shows what is connected to MySQL, which dimensions are available, and which source-quality conditions remain visible in the planning view.', after=8)
    add_table(doc, ['Source or view', 'Verified content'], [
        ['Customer.xlsx', '2,159 customer rows.'],
        ['Item.xlsx', '5,706 item rows.'],
        ['Sale Order History.xlsx', '272,000 sales-order-history lines.'],
        ['RawData_Combined.xlsx', 'Exact copies of the three original sheets.'],
        ['Sales_Order_Merged in MySQL', 'Original sales fields and repaired dates, prices, quantities, derived revenue, and quality flags.'],
        ['Canonical forecast view', '265,135 rows after removing extra exact duplicate copies for modeling.'],
    ], widths=[2.3, 4.5], font_size=9.1)
    add_section_heading(doc, 'Coverage and quality signals', 2)
    for item in [
        'Source coverage runs from April 2021 through March 2025 across 48 monthly periods.',
        'The source contains 87 product categories and 69 industries.',
        'The database preserves the original sales dates from 2021-04-01 through 2025-03-31.',
        '6,865 extra exact duplicate lines are flagged; 11,928 rows belong to duplicate groups.',
        '13 items have missing list prices; 15 sales lines retain null list-price-derived values.',
        'The original Mitsubishi source has no vendor relationship, so vendor-level views require an imported mapping.',
    ]:
        add_bullet(doc, item)
    add_text(doc, 'The first SQL load had stored zero dates and carriage returns in month labels. Those artifacts were repaired from the original source workbooks before the current seed was generated.', after=8)

    # 7
    add_page_break(doc)
    add_section_heading(doc, '7 CSV imports and data quality', 1)
    add_text(doc, 'The Import CSV button is available in the top-right controls. Importing is preview-first: the file is read, normalized, validated, and shown as sample rows before anything is committed.', after=8)
    add_table(doc, ['Import type', 'Required columns', 'Optional columns or behavior'], [
        ['Sales', 'month, item, actual_units, revenue', 'product, customer, industry, vendor; common source aliases are accepted.'],
        ['Vendor mapping', 'item, vendor', 'Each item must map to one vendor.'],
        ['Inventory mapping', 'item plus on_hand_units or lead_time_days', 'Either numeric inventory column may be supplied.'],
    ], widths=[1.45, 2.45, 2.9], font_size=8.9)
    add_section_heading(doc, 'Import safeguards', 2)
    for item in [
        'UTF-8 CSV validation with a 20 MB file limit and a 250,000-row limit.',
        'Date, quantity, revenue, item, vendor, on-hand, and lead-time validation.',
        'Append or replace-active modes for sales overlays.',
        'Exact duplicate detection and idempotent import tokens.',
        'Original SQL source tables are never overwritten.',
        'Imported overlays refresh the forecast and classifier data adapter after commit.',
    ]:
        add_bullet(doc, item)
    add_text(doc, 'Sales CSV template header:', bold=True, after=2)
    add_code(doc, 'month,item,product,customer,industry,vendor,actual_units,revenue')

    # 8
    add_section_heading(doc, '8 Architecture and operations', 1)
    add_table(doc, ['Layer', 'Implementation'], [
        ['Frontend', 'React, TypeScript, and Vite on port 5174.'],
        ['Backend', 'FastAPI, SQLAlchemy, and PyMySQL on port 8100.'],
        ['Database', 'MySQL database named missubibi.'],
        ['Forecast execution', 'Bounded background jobs with isolated model workers.'],
        ['Persistence', 'Local job records, forecast registry, fingerprints, diagnostics, and results.'],
        ['Scheduler', 'Hourly check for changed data or a new UTC month; refreshes requested configurations.'],
        ['Protection', 'Optional PLANNING_API_KEY for API routes.'],
    ], widths=[1.7, 5.1], font_size=9.2)
    add_text(doc, 'Model workers isolate native statistical, machine-learning, and neural-network libraries so one engine failure or native threading conflict does not take down the API process. The scheduler refreshes previously requested configurations; it does not launch exhaustive training for every possible item, customer, industry, and vendor combination.', after=8)
    add_text(doc, 'The model registry versions reproducible forecast runs and active revisions. It stores parameters, data and code fingerprints, diagnostics, and results rather than a complete binary copy of every fitted estimator.', after=8)

    # 9
    add_section_heading(doc, '9 Teammate setup and database synchronization', 1)
    add_text(doc, 'The shared GitHub branch contains the repaired Mitsubishi SQL seed, the dashboard implementation, the startup scripts, and the database freshness check. Teammates should use the working branch rather than the older default branch.', after=8)
    add_code(doc, 'git clone -b codex/forecast-planning-workspace https://github.com/sokolovxf/missubibi.git\ncd missubibi\n./start-local.sh --sync-db')
    add_text(doc, 'For an existing clone:', bold=True, after=2)
    add_code(doc, 'git fetch origin\ngit checkout codex/forecast-planning-workspace\ngit pull\n./start-local.sh --sync-db')
    add_text(doc, 'The startup command compares the checked-out schema.sql and seed.sql fingerprints with a marker recorded inside MySQL. If the local database is already current, it is left untouched. If it is missing or stale, the script rebuilds the local missubibi database from the checked-out files and then starts the backend and frontend.', after=8)
    add_text(doc, 'Windows PowerShell:', bold=True, after=2)
    add_code(doc, '.\\start-local.ps1 -SyncDatabase')
    add_text(doc, 'The database refresh is intentionally explicit because it drops and recreates the local missubibi database. The metadata marker is excluded from future database exports, so it does not become part of the shared source seed.', after=8)
    add_table(doc, ['Service', 'Address'], [
        ['Frontend', 'http://localhost:5174'],
        ['Backend', 'http://127.0.0.1:8100'],
        ['Health check', 'http://127.0.0.1:8100/api/health'],
    ], widths=[2.0, 4.8], font_size=9.4)

    # 10
    add_section_heading(doc, '10 Interpretation rules and operational boundaries', 1)
    add_text(doc, 'These rules keep the dashboard results honest and prevent a planning control from being mistaken for a model guarantee.', after=8)
    for item in [
        'Forecast accuracy is historical held-out evidence, not a guarantee of future accuracy.',
        'Scenario fit is separate from model validation accuracy.',
        'Daily and weekly forecast values may be allocated from the monthly modeling contract, while native transaction-date history is available as a separate view.',
        'Vendor-level planning requires a vendor mapping because the original Mitsubishi source does not include vendor relationships.',
        'Inventory recommendations depend on planner-supplied on-hand units and supplier lead-time values.',
        'Classifier signals can be exposed experimentally, but they are disabled by default in the main forecast service.',
        'Weighted ensemble bounds are approximate planning ranges rather than calibrated statistical intervals.',
        'The current application is a local single-process deployment with local MySQL and local job storage. Optional API-key protection is available through PLANNING_API_KEY.',
    ]:
        add_bullet(doc, item)
    add_section_heading(doc, 'Quick reference', 2)
    add_table(doc, ['Need', 'Use'], [
        ['Forecast a portfolio or slice', 'Forecast workspace → scope → measure → frequency → horizon → model → Run forecast.'],
        ['Compare model evidence', 'Model performance → inspect accuracy, WAPE, MASE, folds, weights, and sparklines.'],
        ['Prioritize item review', 'Demand classifier → train or refresh → select item or all → Check demand → review borderline rows.'],
        ['Plan a replenishment', 'Lead-time inventory → set assumptions → apply → edit on-hand → review suggested order and dates.'],
        ['Verify source coverage', 'Data & coverage → inspect dimensions, date coverage, duplicate flags, missing prices, and fingerprint.'],
        ['Load new operational data', 'Import CSV → preview → validate → append or replace-active → commit.'],
        ['Start with current SQL data', 'Run ./start-local.sh --sync-db or .\\start-local.ps1 -SyncDatabase.'],
    ], widths=[2.25, 4.55], font_size=8.9)
    add_text(doc, 'The complete implementation is available on the GitHub branch codex/forecast-planning-workspace.', color=MUTED, italic=True, after=0)

    doc.core_properties.title = 'Mitsubishi Electric Sales Intelligence System'
    doc.core_properties.subject = 'System overview and operating guide'
    doc.core_properties.author = 'Missubibi'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
